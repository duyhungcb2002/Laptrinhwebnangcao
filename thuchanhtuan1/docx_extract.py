from pathlib import Path
from docx import Document


FILES = [
    Path(r"C:\Users\HN STORE\Downloads\Rubric_Phan_tich_Thiet_ke_HTTT.docx"),
    Path(r"C:\Users\HN STORE\Downloads\Website_Thuong_mai_Dien_tu_Linh_kien_May_tinh.docx"),
]


for path in FILES:
    print("\n" + "=" * 90)
    print(path.name)
    print("=" * 90)
    document = Document(path)

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{index:03d} [{paragraph.style.name}] {text}")

    for table_index, table in enumerate(document.tables):
        print(f"\nTABLE {table_index} ({len(table.rows)}x{len(table.columns)})")
        for row_index, row in enumerate(table.rows):
            values = [" ".join(cell.text.split()) for cell in row.cells]
            print(f"R{row_index:02d}: " + " || ".join(values))
