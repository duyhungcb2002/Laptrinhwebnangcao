from __future__ import annotations

import math
import sys
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(r"D:\HMNU\laptrinhwweb\thuchanhtuan1")
OUTPUT = ROOT / "Tai_lieu_Phan_tich_Thiet_ke_TechHub_PC.docx"
ASSET_DIR = ROOT / "techhub_report_assets"
ASSET_DIR.mkdir(exist_ok=True)

SKILL_DIR = Path(
    r"C:\Users\HN STORE\.codex\plugins\cache\openai-primary-runtime\documents\26.813.12317\skills\documents"
)
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


# Resolved design preset: standard_business_brief.
# Named override: academic report cover via editorial_cover header pattern.
COLORS = {
    "navy": "0B2545",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "tech_blue": "004AC6",
    "light_blue": "E8EEF5",
    "lighter_blue": "F3F7FC",
    "gray": "5B6573",
    "light_gray": "F2F4F7",
    "border": "C9D2DF",
    "green": "137A4A",
    "green_fill": "E8F5ED",
    "orange": "C65D00",
    "orange_fill": "FFF1E6",
    "red": "9B1C1C",
    "red_fill": "FDECEC",
    "white": "FFFFFF",
    "black": "191C1E",
}

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 90, "bottom": 90, "start": 120, "end": 120}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="C9D2DF", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=COLORS["gray"])
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-1" \\h \\z \\u '
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "Mở tài liệu trong Microsoft Word và chọn Update Field để cập nhật mục lục."
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2, fallback, fld_char3])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(COLORS["black"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, COLORS["blue"], 16, 8),
        "Heading 2": (13, COLORS["blue"], 12, 6),
        "Heading 3": (12, COLORS["dark_blue"], 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    if "Caption TechHub" not in styles:
        caption = styles.add_style("Caption TechHub", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption TechHub"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(COLORS["gray"])
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("TECHHUB PC  |  PHÂN TÍCH & THIẾT KẾ HỆ THỐNG THÔNG TIN")
    set_run_font(hr, size=8.5, color=COLORS["gray"], bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    add_page_field(fp)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_para(doc, text="", *, bold_lead=None, italic=False, align=None, color=None, after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, bold=True, color=color or COLORS["black"])
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, italic=italic, color=color or COLORS["black"])
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color or COLORS["black"])
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    return p


def add_number(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_callout(doc, label, text, fill="E8EEF5", color="0B2545"):
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, color=fill, size="2")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=color)
    apply_table_geometry(
        table,
        [CONTENT_WIDTH_DXA],
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=180,
        cell_margins_dxa={"top": 130, "bottom": 130, "start": 180, "end": 180},
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, weights, *, font_size=8.5, header_fill=None, zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_fill = header_fill or COLORS["light_blue"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, header_fill)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=COLORS["navy"])
    set_repeat_table_header(table.rows[0])
    set_table_row_no_split(table.rows[0])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_table_row_no_split(table.rows[-1])
        for idx, value in enumerate(values):
            cell = cells[idx]
            if zebra and row_index % 2 == 1:
                shade_cell(cell, "F8FAFC")
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=COLORS["black"])

    widths = column_widths_from_weights(weights, CONTENT_WIDTH_DXA)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, image_path: Path, caption: str, width=6.25, max_height=7.35):
    with Image.open(image_path) as source_image:
        image_width, image_height = source_image.size
    width = min(width, max_height * image_width / image_height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split(".", 1)[0])
    c = doc.add_paragraph(caption, style="Caption TechHub")
    return c


FONT_REGULAR = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def fnt(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def wrap_text(draw, text, font, max_width):
    words = str(text).split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else current + " " + word
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def rounded_box(draw, xy, text, *, fill="#FFFFFF", outline="#2E74B5", font=None,
                text_fill="#0B2545", radius=18, width=3, padding=18, align="center"):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    font = font or fnt(26)
    lines = wrap_text(draw, text, font, (x2 - x1) - 2 * padding)
    line_height = font.size + 7
    total_h = len(lines) * line_height
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        if align == "left":
            x = x1 + padding
        else:
            x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=text_fill)
        y += line_height


def draw_arrow(draw, start, end, color="#2E74B5", width=5, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        steps = 18
        for i in range(0, steps, 2):
            a = i / steps
            b = min((i + 1) / steps, 1)
            draw.line((x1 + (x2 - x1) * a, y1 + (y2 - y1) * a,
                       x1 + (x2 - x1) * b, y1 + (y2 - y1) * b), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 16
    points = [
        (x2, y2),
        (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6)),
        (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=color)


def draw_actor(draw, center, label, color="#0B2545"):
    x, y = center
    draw.ellipse((x - 24, y - 95, x + 24, y - 47), outline=color, width=5)
    draw.line((x, y - 47, x, y + 25), fill=color, width=5)
    draw.line((x - 40, y - 15, x + 40, y - 15), fill=color, width=5)
    draw.line((x, y + 25, x - 36, y + 82), fill=color, width=5)
    draw.line((x, y + 25, x + 36, y + 82), fill=color, width=5)
    bbox = draw.textbbox((0, 0), label, font=fnt(28, True))
    draw.text((x - (bbox[2] - bbox[0]) / 2, y + 95), label, font=fnt(28, True), fill=color)


def diagram_canvas(size, title):
    img = Image.new("RGB", size, "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((50, 30), title, font=fnt(38, True), fill="#0B2545")
    draw.line((50, 90, size[0] - 50, 90), fill="#C9D2DF", width=3)
    return img, draw


def make_use_case_diagram(path):
    img, draw = diagram_canvas((1800, 1120), "Use Case tổng thể - TechHub PC")
    draw.rounded_rectangle((300, 125, 1500, 1040), radius=24, fill="#F8FAFC", outline="#004AC6", width=4)
    draw.text((330, 145), "Biên hệ thống TechHub PC", font=fnt(28, True), fill="#004AC6")
    use_cases = [
        ("Xem / tìm / lọc sản phẩm", (450, 235, 820, 330)),
        ("Đăng ký / đăng nhập", (980, 235, 1350, 330)),
        ("Quản lý giỏ hàng", (450, 400, 820, 495)),
        ("Checkout / đặt hàng", (980, 400, 1350, 495)),
        ("Theo dõi / hủy đơn", (450, 565, 820, 660)),
        ("Đánh giá sản phẩm", (980, 565, 1350, 660)),
        ("Quản lý sản phẩm / danh mục", (450, 730, 820, 825)),
        ("Quản lý kho / đơn hàng", (980, 730, 1350, 825)),
        ("Quản lý người dùng / thống kê", (715, 895, 1085, 990)),
    ]
    for text, xy in use_cases:
        rounded_box(draw, xy, text, fill="#FFFFFF", outline="#2E74B5", font=fnt(23), radius=42)
    draw_actor(draw, (150, 310), "Guest")
    draw_actor(draw, (150, 625), "Customer")
    draw_actor(draw, (1650, 650), "Admin")
    for end in [(450, 280), (980, 280)]:
        draw_arrow(draw, (205, 310), end, width=3)
    for end in [(450, 447), (980, 447), (450, 612), (980, 612)]:
        draw_arrow(draw, (205, 625), end, width=3)
    for end in [(1350, 777), (1350, 777), (1085, 942)]:
        draw_arrow(draw, (1595, 650), end, width=3)
    img.save(path)


def make_context_diagram(path):
    img, draw = diagram_canvas((1800, 920), "System Context - phạm vi và tác nhân")
    rounded_box(draw, (620, 260, 1180, 650),
                "TECHHUB PC\nWebsite thương mại điện tử linh kiện máy tính\n\nQuản lý catalog, giỏ hàng, đơn hàng, kho, đánh giá và báo cáo",
                fill="#E8EEF5", outline="#004AC6", font=fnt(27), radius=28)
    boxes = [
        ((70, 210, 430, 370), "Guest\nXem và tìm sản phẩm"),
        ((70, 515, 430, 675), "Customer\nMua hàng và theo dõi đơn"),
        ((1370, 210, 1730, 370), "Admin\nVận hành hệ thống"),
        ((1370, 515, 1730, 675), "Dịch vụ thanh toán giả lập\nXác nhận giao dịch mô phỏng"),
    ]
    for xy, text in boxes:
        rounded_box(draw, xy, text, fill="#FFFFFF", outline="#2E74B5", font=fnt(24))
    draw_arrow(draw, (430, 290), (620, 360), width=5)
    draw_arrow(draw, (430, 595), (620, 540), width=5)
    draw_arrow(draw, (1370, 290), (1180, 360), width=5)
    draw_arrow(draw, (1180, 540), (1370, 595), width=5)
    draw.text((670, 735), "HTTPS/JSON • JWT/RBAC • Default deny", font=fnt(25, True), fill="#137A4A")
    img.save(path)


def make_container_diagram(path):
    img, draw = diagram_canvas((1800, 1080), "Container / Module - kiến trúc mục tiêu")
    rounded_box(draw, (80, 170, 430, 340), "Trình duyệt\nReact 19 + Vite\nTailwind + Router + Axios",
                fill="#F3F7FC", outline="#004AC6", font=fnt(24))
    rounded_box(draw, (610, 145, 1190, 370), "ASP.NET Core Web API (.NET 10)\n\nControllers • Application Services\nAuthorization Policies • EF Core\nProblemDetails • Health Checks",
                fill="#E8EEF5", outline="#004AC6", font=fnt(25))
    rounded_box(draw, (1370, 170, 1720, 340), "PostgreSQL\nDữ liệu nghiệp vụ\nRBAC • Audit log",
                fill="#F3F7FC", outline="#2E74B5", font=fnt(25))
    draw_arrow(draw, (430, 255), (610, 255), width=5)
    draw.text((462, 215), "HTTPS / JSON", font=fnt(20, True), fill="#5B6573")
    draw_arrow(draw, (1190, 255), (1370, 255), width=5)
    draw.text((1230, 215), "EF Core", font=fnt(20, True), fill="#5B6573")

    modules = [
        ("Identity & RBAC", (105, 520, 415, 660)),
        ("Catalog", (455, 520, 765, 660)),
        ("Cart & Checkout", (805, 520, 1115, 660)),
        ("Orders", (1155, 520, 1465, 660)),
        ("Inventory", (280, 745, 590, 885)),
        ("Reviews", (745, 745, 1055, 885)),
        ("Reporting", (1210, 745, 1520, 885)),
    ]
    for text, xy in modules:
        rounded_box(draw, xy, text, fill="#FFFFFF", outline="#7EA4D4", font=fnt(23))
    draw.text((80, 450), "Các module nghiệp vụ trong modular monolith", font=fnt(28, True), fill="#0B2545")
    draw.text((90, 980), "Ranh giới module rõ; triển khai một API để phù hợp phạm vi 6 tuần.", font=fnt(24), fill="#5B6573")
    img.save(path)


def entity_box(draw, xy, name, fields):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill="#FFFFFF", outline="#2E74B5", width=3)
    draw.rectangle((x1, y1, x2, y1 + 48), fill="#E8EEF5")
    draw.text((x1 + 12, y1 + 9), name, font=fnt(22, True), fill="#0B2545")
    y = y1 + 58
    for field in fields:
        draw.text((x1 + 12, y), field, font=fnt(16), fill="#191C1E")
        y += 25


def make_erd(path):
    img, draw = diagram_canvas((2200, 1650), "ERD rút gọn - khóa và quan hệ cốt lõi")
    entities = {
        "User": ((50, 150, 360, 335), ["PK user_id", "UQ email", "status", "password_hash"]),
        "Role": ((500, 150, 800, 310), ["PK role_id", "UQ name"]),
        "Permission": ((950, 150, 1300, 310), ["PK permission_id", "UQ code"]),
        "UserRole": ((420, 390, 740, 550), ["PK/FK user_id", "PK/FK role_id"]),
        "RolePermission": ((860, 390, 1250, 550), ["PK/FK role_id", "PK/FK permission_id"]),
        "Category": ((1450, 150, 1780, 335), ["PK category_id", "UQ slug", "name", "is_active"]),
        "Product": ((1450, 430, 1800, 665), ["PK product_id", "FK category_id", "UQ sku", "price", "stock", "row_version"]),
        "InventoryTxn": ((1830, 760, 2160, 970), ["PK txn_id", "FK product_id", "type", "quantity", "created_by"]),
        "Cart": ((50, 760, 340, 920), ["PK cart_id", "FK user_id"]),
        "CartItem": ((430, 760, 760, 970), ["PK cart_item_id", "FK cart_id", "FK product_id", "quantity"]),
        "Order": ((50, 1120, 380, 1380), ["PK order_id", "FK user_id", "status", "total_amount", "shipping_snapshot", "created_at"]),
        "OrderItem": ((520, 1120, 870, 1380), ["PK order_item_id", "FK order_id", "FK product_id", "unit_price_snapshot", "quantity"]),
        "Payment": ((1030, 1120, 1360, 1330), ["PK payment_id", "FK order_id", "method", "status", "transaction_ref"]),
        "Review": ((1510, 1120, 1840, 1380), ["PK review_id", "FK user_id", "FK product_id", "FK order_item_id", "rating 1..5", "status"]),
    }
    for name, (xy, fields) in entities.items():
        entity_box(draw, xy, name, fields)
    relations = [
        ((360, 245), (500, 470), "1..N"),
        ((740, 470), (650, 310), "N..1"),
        ((800, 235), (860, 470), "1..N"),
        ((1250, 470), (1130, 310), "N..1"),
        ((1780, 245), (1625, 430), "1..N"),
        ((1800, 570), (1990, 760), "1..N"),
        ((205, 335), (195, 760), "1..1"),
        ((340, 840), (430, 850), "1..N"),
        ((760, 860), (1450, 570), "N..1"),
        ((205, 335), (215, 1120), "1..N"),
        ((380, 1250), (520, 1250), "1..N"),
        ((870, 1230), (1450, 600), "N..1"),
        ((380, 1310), (1030, 1230), "1..N"),
        ((1840, 1230), (1650, 665), "N..1"),
        ((1510, 1280), (360, 300), "N..1"),
    ]
    for start, end, label in relations:
        draw.line((*start, *end), fill="#7A8AA0", width=3)
        mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
        draw.text((mx + 5, my - 20), label, font=fnt(15, True), fill="#5B6573")
    draw.text((50, 1530), "PK: khóa chính • FK: khóa ngoại • UQ: duy nhất • Giá và địa chỉ được snapshot tại thời điểm đặt hàng",
              font=fnt(22), fill="#5B6573")
    img.save(path)


def make_sequence(path):
    img, draw = diagram_canvas((1900, 1320), "Sequence Diagram - tạo đơn hàng")
    actors = [(160, "Customer"), (520, "React SPA"), (900, ".NET API"), (1280, "Order Service"), (1660, "PostgreSQL")]
    for x, label in actors:
        rounded_box(draw, (x - 125, 140, x + 125, 220), label, fill="#F3F7FC", outline="#2E74B5", font=fnt(21))
        draw.line((x, 220, x, 1220), fill="#9AA7B8", width=3)
    events = [
        (280, 160, 520, "1. Chọn Checkout"),
        (370, 520, 900, "2. POST /api/orders + JWT"),
        (465, 900, 1280, "3. Kiểm tra quyền và dữ liệu"),
        (560, 1280, 1660, "4. BEGIN; khóa các dòng sản phẩm"),
        (655, 1660, 1280, "5. Giá + tồn kho hiện tại"),
        (750, 1280, 1660, "6. Tạo Order/OrderItems; trừ kho"),
        (845, 1660, 1280, "7. COMMIT"),
        (940, 1280, 900, "8. OrderCreated"),
        (1035, 900, 520, "9. 201 Created + orderId"),
        (1130, 520, 160, "10. Hiển thị đặt hàng thành công"),
    ]
    for y, x1, x2, label in events:
        color = "#137A4A" if x2 < x1 else "#2E74B5"
        draw_arrow(draw, (x1, y), (x2, y), color=color, width=4)
        draw.text((min(x1, x2) + 12, y - 34), label, font=fnt(19), fill="#191C1E")
    draw.rounded_rectangle((1210, 515, 1730, 900), radius=20, outline="#C65D00", width=4)
    draw.text((1240, 875), "Cùng một transaction", font=fnt(18, True), fill="#C65D00")
    img.save(path)


def make_activity(path, inventory=False):
    title = "Activity Diagram - nhập kho" if inventory else "Activity Diagram - đặt hàng"
    img, draw = diagram_canvas((1500, 1650), title)
    cx = 750
    draw.ellipse((cx - 22, 125, cx + 22, 169), fill="#0B2545")
    if inventory:
        steps = [
            (210, "Admin mở chức năng Nhập kho", "normal"),
            (355, "Nhập sản phẩm, số lượng và ghi chú", "normal"),
            (500, "Dữ liệu hợp lệ?", "decision"),
            (680, "Ghi InventoryTransaction = IMPORT", "normal"),
            (825, "Cộng tồn kho trong transaction", "normal"),
            (970, "Ghi audit log và commit", "normal"),
            (1115, "Trả kết quả và số tồn mới", "normal"),
        ]
    else:
        steps = [
            (210, "Customer gửi yêu cầu đặt hàng", "normal"),
            (355, "Xác thực JWT và quyền order.create", "normal"),
            (500, "Thông tin nhận hàng hợp lệ?", "decision"),
            (680, "Khóa và đọc giá/tồn kho hiện tại", "normal"),
            (825, "Đủ hàng và giá còn hiệu lực?", "decision"),
            (1005, "Tạo đơn + chi tiết; snapshot giá", "normal"),
            (1150, "Trừ tồn kho và commit", "normal"),
            (1295, "Trả 201 Created", "normal"),
        ]
    previous_y = 169
    for y, text, kind in steps:
        draw_arrow(draw, (cx, previous_y), (cx, y - 55), width=4)
        if kind == "decision":
            pts = [(cx, y - 60), (cx + 210, y), (cx, y + 60), (cx - 210, y)]
            draw.polygon(pts, fill="#FFF1E6", outline="#C65D00")
            lines = wrap_text(draw, text, fnt(21, True), 300)
            ty = y - len(lines) * 14
            for line in lines:
                bbox = draw.textbbox((0, 0), line, font=fnt(21, True))
                draw.text((cx - (bbox[2] - bbox[0]) / 2, ty), line, font=fnt(21, True), fill="#7A3D00")
                ty += 28
            draw_arrow(draw, (cx + 210, y), (1250, y), color="#9B1C1C", width=4)
            rounded_box(draw, (1160, y - 55, 1440, y + 55), "Trả lỗi 400/409", fill="#FDECEC", outline="#9B1C1C", font=fnt(18))
            draw.text((1000, y - 32), "Không", font=fnt(17, True), fill="#9B1C1C")
            draw.text((cx + 20, y + 60), "Có", font=fnt(17, True), fill="#137A4A")
            previous_y = y + 60
        else:
            rounded_box(draw, (cx - 300, y - 55, cx + 300, y + 55), text,
                        fill="#F3F7FC", outline="#2E74B5", font=fnt(21))
            previous_y = y + 55
    draw_arrow(draw, (cx, previous_y), (cx, 1495), width=4)
    draw.ellipse((cx - 27, 1495, cx + 27, 1549), outline="#0B2545", width=5)
    draw.ellipse((cx - 18, 1504, cx + 18, 1540), fill="#0B2545")
    img.save(path)


def make_trust_diagram(path):
    img, draw = diagram_canvas((1800, 1050), "Trust boundaries và entry points")
    zones = [
        ((50, 140, 520, 930), "Thiết bị người dùng", "#F8FAFC"),
        ((570, 140, 1180, 930), "Vùng ứng dụng", "#F3F7FC"),
        ((1230, 140, 1750, 930), "Vùng dữ liệu", "#F8FAFC"),
    ]
    for xy, label, fill in zones:
        draw.rounded_rectangle(xy, radius=22, fill=fill, outline="#9B1C1C", width=4)
        draw.text((xy[0] + 20, xy[1] + 18), label, font=fnt(24, True), fill="#9B1C1C")
    rounded_box(draw, (125, 300, 445, 470), "Browser\nEntry: form, URL, file upload", fill="#FFFFFF", outline="#2E74B5", font=fnt(22))
    rounded_box(draw, (680, 260, 1070, 450), "React SPA\nKhông phải security boundary\nKhông tin dữ liệu client", fill="#FFFFFF", outline="#2E74B5", font=fnt(22))
    rounded_box(draw, (680, 600, 1070, 790), ".NET Web API\nAuthN + AuthZ + validation\nRate limit + audit", fill="#E8EEF5", outline="#004AC6", font=fnt(22))
    rounded_box(draw, (1325, 300, 1655, 470), "PostgreSQL\nPII, đơn hàng, kho, RBAC", fill="#FFFFFF", outline="#2E74B5", font=fnt(22))
    rounded_box(draw, (1325, 650, 1655, 800), "Media storage\nẢnh sản phẩm", fill="#FFFFFF", outline="#2E74B5", font=fnt(22))
    draw_arrow(draw, (445, 385), (680, 355), width=5)
    draw.text((490, 320), "HTTPS", font=fnt(20, True), fill="#137A4A")
    draw_arrow(draw, (875, 450), (875, 600), width=5)
    draw.text((900, 505), "JSON/JWT", font=fnt(18, True), fill="#5B6573")
    draw_arrow(draw, (1070, 690), (1325, 385), width=5)
    draw_arrow(draw, (1070, 720), (1325, 725), width=5)
    draw.text((80, 965), "Ranh giới đỏ: thay đổi mức tin cậy. Mọi yêu cầu qua ranh giới phải được xác thực, phân quyền và kiểm tra dữ liệu.",
              font=fnt(21), fill="#5B6573")
    img.save(path)


def make_jwt_flow(path):
    img, draw = diagram_canvas((1800, 950), "JWT / Refresh Token flow")
    actors = [(170, "Browser"), (650, "Auth API"), (1130, "Identity/RBAC"), (1620, "Database")]
    for x, label in actors:
        rounded_box(draw, (x - 120, 130, x + 120, 205), label, fill="#F3F7FC", outline="#2E74B5", font=fnt(21))
        draw.line((x, 205, x, 870), fill="#9AA7B8", width=3)
    events = [
        (270, 170, 650, "POST /login"),
        (360, 650, 1130, "Xác minh mật khẩu + trạng thái"),
        (450, 1130, 1620, "Đọc roles/permissions"),
        (540, 1620, 1130, "User + quyền"),
        (630, 1130, 650, "Claims: sub, role, permission, jti"),
        (720, 650, 170, "Access JWT 15 phút + refresh cookie 7 ngày"),
        (810, 170, 650, "Authorization: Bearer <access token>"),
    ]
    for y, x1, x2, label in events:
        draw_arrow(draw, (x1, y), (x2, y), color="#137A4A" if x2 < x1 else "#2E74B5", width=4)
        draw.text((min(x1, x2) + 12, y - 32), label, font=fnt(18), fill="#191C1E")
    img.save(path)


def generate_diagrams():
    paths = {
        "use_case": ASSET_DIR / "use_case.png",
        "context": ASSET_DIR / "system_context.png",
        "container": ASSET_DIR / "container_module.png",
        "erd": ASSET_DIR / "erd.png",
        "sequence": ASSET_DIR / "sequence_order.png",
        "activity_order": ASSET_DIR / "activity_order.png",
        "activity_inventory": ASSET_DIR / "activity_inventory.png",
        "trust": ASSET_DIR / "trust_boundaries.png",
        "jwt": ASSET_DIR / "jwt_flow.png",
    }
    make_use_case_diagram(paths["use_case"])
    make_context_diagram(paths["context"])
    make_container_diagram(paths["container"])
    make_erd(paths["erd"])
    make_sequence(paths["sequence"])
    make_activity(paths["activity_order"], inventory=False)
    make_activity(paths["activity_inventory"], inventory=True)
    make_trust_diagram(paths["trust"])
    make_jwt_flow(paths["jwt"])
    return paths


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO PHÂN TÍCH & THIẾT KẾ\nHỆ THỐNG THÔNG TIN")
    set_run_font(r, size=13, bold=True, color=COLORS["gray"])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("TECHHUB PC")
    set_run_font(r, size=32, bold=True, color=COLORS["navy"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Website thương mại điện tử kinh doanh linh kiện máy tính")
    set_run_font(r, size=15, bold=True, color=COLORS["blue"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(62)
    r = p.add_run("Use Case • User Story • Requirements • Database Domain\nSystem Design • RBAC/JWT • Threat Model • Traceability")
    set_run_font(r, size=10.5, italic=True, color=COLORS["gray"])

    for label in ("Lớp", "Nhóm/Sinh viên", "Giảng viên"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(f"{label}:  ................................................................................")
        set_run_font(r, size=11, color=COLORS["black"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    r = p.add_run("Hà Nội, tháng 08 năm 2026")
    set_run_font(r, size=11, bold=True, color=COLORS["gray"])
    doc.add_page_break()


def add_document_control(doc):
    add_heading(doc, "THÔNG TIN KIỂM SOÁT TÀI LIỆU", 1)
    add_table(
        doc,
        ["Thuộc tính", "Nội dung"],
        [
            ("Tên hệ thống", "TechHub PC - Website thương mại điện tử linh kiện máy tính"),
            ("Loại tài liệu", "Hồ sơ phân tích và thiết kế hệ thống thông tin"),
            ("Phiên bản", "1.0 - Thiết kế mục tiêu cho phiên bản đầu"),
            ("Ngày lập", "21/08/2026"),
            ("Trạng thái", "Dùng để review và triển khai; cần bổ sung thông tin lớp/nhóm trước khi nộp"),
            ("Quy tắc tên file nộp", "TenLop_TenNhom_TenDeTai.docx"),
        ],
        [1.5, 5.0],
        font_size=9.5,
    )
    add_callout(
        doc,
        "Trạng thái triển khai",
        "Frontend đã hoàn thành nền Vite/React, Router, Tailwind và Axios. Các module nghiệp vụ, backend .NET 10 và PostgreSQL trong tài liệu là thiết kế mục tiêu sẽ được triển khai theo lộ trình; tài liệu không tuyên bố các phần đó đã hoàn thành.",
        fill=COLORS["orange_fill"],
        color="7A3D00",
    )
    doc.add_page_break()
    add_heading(doc, "MỤC LỤC", 1)
    p = doc.add_paragraph()
    add_toc_field(p)
    doc.add_page_break()


def add_summary(doc):
    add_heading(doc, "TÓM TẮT ĐIỀU HÀNH", 1)
    add_para(doc, "TechHub PC là hệ thống thương mại điện tử chuyên kinh doanh linh kiện máy tính. Hệ thống hỗ trợ ba nhóm tác nhân: Guest, Customer và Admin; bao phủ từ tìm kiếm sản phẩm, giỏ hàng, checkout đến quản lý kho, đơn hàng, đánh giá và thống kê.")
    add_para(doc, "Thiết kế chọn kiến trúc modular monolith gồm React SPA, ASP.NET Core Web API .NET 10 và PostgreSQL. Cách tiếp cận này giữ chi phí triển khai phù hợp kế hoạch học kỳ, đồng thời duy trì ranh giới module để có thể mở rộng khi tải hoặc đội phát triển tăng.")
    add_callout(doc, "Mục tiêu thiết kế", "Tạo một hồ sơ đủ chi tiết để người phát triển có thể triển khai mà không phải hỏi lại các quyết định quan trọng; mọi luồng ưu tiên đều truy vết được từ mục tiêu, yêu cầu, dữ liệu, API đến kiểm thử.")
    add_heading(doc, "Cấu trúc hồ sơ theo rubric", 2)
    add_table(
        doc,
        ["Nhóm chấm", "Phần trong tài liệu", "Bằng chứng chính"],
        [
            ("Phân tích & thiết kế", "Chương 1-8", "Goal, Use Case, User Story/AC, FR/NFR, ERD, context, module, sequence, API"),
            ("Trade-off & ADR", "Chương 9", "Bốn quyết định có options, rationale, consequences và trigger xem xét lại"),
            ("Security", "Chương 10", "Trust boundary, threat rating, RBAC matrix, JWT flow, negative tests"),
            ("Tài liệu & truy vết", "Chương 11-12", "Test catalog, RTM Goal→UC→FR/NFR→Entity/API→Test"),
            ("Mở rộng & vận hành", "Chương 13", "Tải giả định, index/cache/pagination, health/log/audit, debt và roadmap"),
        ],
        [1.35, 1.25, 3.9],
        font_size=8.5,
    )


def add_context_scope(doc, diagrams):
    add_heading(doc, "1. BỐI CẢNH, MỤC TIÊU VÀ PHẠM VI", 1)
    add_heading(doc, "1.1. Bối cảnh và vấn đề", 2)
    add_para(doc, "Cửa hàng linh kiện máy tính cần một kênh bán hàng trực tuyến tập trung. Nếu quản lý sản phẩm, đơn hàng và tồn kho bằng các bảng rời, dữ liệu dễ lệch, khó truy nguồn xuất/nhập và có nguy cơ nhận đơn vượt tồn kho. TechHub PC giải quyết vấn đề này bằng một luồng xuyên suốt từ catalog đến vận hành sau bán.")

    add_heading(doc, "1.2. Mục tiêu nghiệp vụ", 2)
    goals = [
        ("G-01", "Khách tìm được sản phẩm phù hợp theo từ khóa, danh mục, giá và tiêu chí sắp xếp."),
        ("G-02", "Customer đặt hàng an toàn; giá và tồn kho được backend xác nhận tại thời điểm tạo đơn."),
        ("G-03", "Customer theo dõi lịch sử, trạng thái và chỉ đánh giá sản phẩm đã mua trong đơn hoàn thành."),
        ("G-04", "Admin quản lý danh mục, sản phẩm, người dùng, kho và đơn hàng theo đúng quyền."),
        ("G-05", "Dữ liệu kho có lịch sử biến động, cảnh báo tồn thấp và không bị âm do cạnh tranh đồng thời."),
        ("G-06", "Quản trị viên xem được doanh thu, đơn hàng, sản phẩm bán chạy và hàng sắp hết theo thời gian."),
    ]
    add_table(doc, ["ID", "Mục tiêu"], goals, [0.75, 5.75], font_size=9.5)

    add_heading(doc, "1.3. Phạm vi phiên bản đầu", 2)
    add_table(
        doc,
        ["Trong phạm vi", "Ngoài phạm vi phiên bản đầu"],
        [
            ("Tài khoản Customer/Admin và RBAC permission-based", "Build PC và kiểm tra tương thích linh kiện"),
            ("Catalog, tìm kiếm, lọc, sắp xếp, chi tiết sản phẩm", "Chatbot/AI tư vấn và gợi ý thông minh"),
            ("Giỏ hàng, checkout, COD và thanh toán giả lập", "Voucher/khuyến mãi phức tạp"),
            ("Đơn hàng, trạng thái, hủy theo điều kiện", "Thanh toán trực tuyến bằng tiền thật"),
            ("Kho, nhập/xuất, lịch sử, cảnh báo tồn thấp", "Tích hợp hãng vận chuyển thực tế"),
            ("Đánh giá đã xác minh và thống kê quản trị", "Microservices, message broker và search engine riêng"),
        ],
        [3.25, 3.25],
        font_size=9,
    )

    add_heading(doc, "1.4. Giả định và ràng buộc", 2)
    for item in [
        "Kế hoạch triển khai cốt lõi trong khoảng 6 tuần bởi một nhóm sinh viên nhỏ.",
        "Thanh toán online được mô phỏng; COD là phương thức chính của phiên bản đầu.",
        "Sản phẩm có một mức tồn có thể bán; chưa tách nhiều kho vật lý.",
        "Frontend và API có thể chạy khác origin trong môi trường phát triển; CORS chỉ cho phép origin đã cấu hình.",
        "Mọi giá trị tiền dùng decimal tại backend và numeric/decimal trong PostgreSQL; không dùng float.",
    ]:
        add_bullet(doc, item)
    add_figure(doc, diagrams["context"], "Hình 1. System Context và phạm vi tương tác của TechHub PC")


def add_actors_usecases(doc, diagrams):
    add_heading(doc, "2. TÁC NHÂN VÀ USE CASE", 1)
    add_heading(doc, "2.1. Tác nhân", 2)
    add_table(
        doc,
        ["Actor", "Mục tiêu", "Quyền tiêu biểu"],
        [
            ("Guest", "Khám phá catalog và trở thành khách hàng", "product.read, category.read, review.read, auth.register, auth.login"),
            ("Customer", "Mua hàng và quản lý thông tin của chính mình", "cart.*, order.create, order.self.read, order.self.cancel, review.create, profile.self.update"),
            ("Admin", "Vận hành cửa hàng", "user.manage, category.*, product.*, inventory.*, order.manage, review.moderate, report.read"),
        ],
        [1.0, 2.0, 3.5],
        font_size=8.8,
    )
    add_figure(doc, diagrams["use_case"], "Hình 2. Use Case Diagram tổng thể và biên hệ thống")

    add_heading(doc, "2.2. Danh mục Use Case", 2)
    use_cases = [
        ("UC-01", "Tra cứu sản phẩm", "Guest/Customer", "Tìm kiếm, lọc, sắp xếp và xem chi tiết"),
        ("UC-02", "Đăng ký/đăng nhập", "Guest", "Tạo phiên xác thực và nạp quyền"),
        ("UC-03", "Quản lý giỏ hàng", "Customer", "Thêm/xóa/đổi số lượng, xem tổng tạm tính"),
        ("UC-04", "Checkout và đặt hàng", "Customer", "Kiểm tra giá/tồn, tạo Order và OrderItem"),
        ("UC-05", "Theo dõi/hủy đơn", "Customer", "Chỉ thao tác trên đơn của chính mình"),
        ("UC-06", "Đánh giá sản phẩm", "Customer", "Chỉ review sản phẩm đã mua và hoàn thành"),
        ("UC-07", "Quản lý sản phẩm", "Admin", "CRUD/ẩn sản phẩm và thông số"),
        ("UC-08", "Quản lý kho", "Admin", "Nhập/xuất, lịch sử và cảnh báo"),
        ("UC-09", "Xử lý đơn hàng", "Admin", "Cập nhật trạng thái theo state machine"),
        ("UC-10", "Quản lý người dùng", "Admin", "Khóa/mở và gán role"),
        ("UC-11", "Xem thống kê", "Admin", "Doanh thu, đơn, top sản phẩm, tồn thấp"),
    ]
    add_table(doc, ["ID", "Tên", "Actor", "Kết quả"], use_cases, [0.65, 1.55, 1.15, 3.15], font_size=8.2)

    detailed = [
        ("2.3. UC-01 - Tra cứu sản phẩm", "Guest hoặc Customer cần tìm sản phẩm phù hợp.",
         ["Người dùng nhập từ khóa hoặc chọn danh mục/bộ lọc.", "Frontend gửi GET /api/products với query và phân trang.", "API kiểm tra giới hạn page/pageSize, truy vấn sản phẩm đang hiển thị và trả tổng số kết quả.", "Người dùng mở trang chi tiết để xem ảnh, giá, tồn và thông số."],
         ["Không có kết quả: trả danh sách rỗng và gợi ý bỏ bớt bộ lọc.", "Tham số không hợp lệ: 400 ProblemDetails.", "Sản phẩm đã ẩn: 404 hoặc không xuất hiện trong danh sách công khai."]),
        ("2.4. UC-04 - Checkout và đặt hàng", "Customer đã đăng nhập, giỏ hàng có ít nhất một sản phẩm.",
         ["Customer nhập người nhận, số điện thoại, địa chỉ, ghi chú và phương thức COD/giả lập.", "Frontend gửi POST /api/orders cùng các productId và quantity.", "Backend xác thực quyền order.create; đọc lại sản phẩm, giá và tồn kho trong transaction.", "Backend tạo Order/OrderItems, snapshot giá/địa chỉ, trừ tồn và ghi biến động kho.", "API trả 201 Created và mã đơn; frontend hiển thị xác nhận."],
         ["JWT thiếu/hết hạn: 401.", "Không có permission: 403.", "Giá thay đổi hoặc không đủ tồn: 409, không tạo đơn và không trừ kho.", "Một sản phẩm bị ẩn: 409 và yêu cầu cập nhật giỏ."]),
        ("2.5. UC-08 - Nhập kho", "Admin có permission inventory.import.",
         ["Admin chọn sản phẩm, nhập số lượng dương và ghi chú/chứng từ.", "API kiểm tra sản phẩm và dữ liệu.", "Trong transaction, hệ thống tạo InventoryTransaction loại IMPORT và cộng tồn.", "Hệ thống ghi AuditLog và trả số tồn mới."],
         ["Không đủ quyền: 403.", "Sản phẩm không tồn tại: 404.", "Số lượng bằng 0/âm: 400.", "Lỗi transaction: rollback toàn bộ."]),
        ("2.6. UC-09 - Cập nhật trạng thái đơn", "Admin có permission order.manage và đơn tồn tại.",
         ["Admin mở chi tiết đơn.", "Chọn trạng thái kế tiếp hợp lệ.", "API kiểm tra state transition, cập nhật và ghi audit.", "Customer nhìn thấy trạng thái mới trong lịch sử đơn."],
         ["Chuyển trạng thái ngược hoặc bỏ qua bước không cho phép: 409.", "Đơn của trạng thái cuối Completed/Cancelled không được sửa: 409."]),
    ]
    for title, precondition, main, alternatives in detailed:
        add_heading(doc, title, 2)
        add_para(doc, f"Tiền điều kiện: {precondition}", bold_lead="Tiền điều kiện:")
        add_para(doc, "Luồng chính:", bold_lead="Luồng chính:")
        for step in main:
            add_number(doc, step)
        add_para(doc, "Luồng thay thế/lỗi:", bold_lead="Luồng thay thế/lỗi:")
        for step in alternatives:
            add_bullet(doc, step)


def add_user_stories(doc):
    add_heading(doc, "3. USER STORY VÀ ACCEPTANCE CRITERIA", 1)
    stories = [
        ("US-01", "Guest", "tìm và lọc sản phẩm", "nhanh chóng chọn linh kiện phù hợp", "UC-01"),
        ("US-02", "Customer", "đặt hàng từ giỏ", "hoàn tất mua hàng với giá và tồn chính xác", "UC-04"),
        ("US-03", "Customer", "xem đơn của mình", "theo dõi quá trình xử lý", "UC-05"),
        ("US-04", "Customer", "hủy đơn Pending", "dừng đơn trước khi cửa hàng xử lý", "UC-05"),
        ("US-05", "Customer", "đánh giá sản phẩm đã mua", "chia sẻ trải nghiệm đáng tin cậy", "UC-06"),
        ("US-06", "Admin", "quản lý sản phẩm và danh mục", "catalog luôn chính xác", "UC-07"),
        ("US-07", "Admin", "nhập/xuất và xem lịch sử kho", "biết nguyên nhân mỗi biến động tồn", "UC-08"),
        ("US-08", "Admin", "xử lý trạng thái đơn", "vận hành đúng quy trình", "UC-09"),
        ("US-09", "Admin", "khóa tài khoản vi phạm", "ngăn truy cập nhưng vẫn bảo toàn lịch sử", "UC-10"),
        ("US-10", "Admin", "xem báo cáo theo thời gian", "đánh giá hoạt động kinh doanh", "UC-11"),
    ]
    add_table(doc, ["ID", "Với vai trò", "Tôi muốn", "Để", "UC"], stories, [0.55, 0.85, 1.65, 2.95, 0.5], font_size=7.8)

    add_heading(doc, "3.1. Acceptance Criteria ưu tiên", 2)
    criteria = [
        ("AC-01", "US-01", "Given catalog có sản phẩm đang hiển thị; When tìm theo tên/SKU và lọc danh mục; Then kết quả đúng, có total và phân trang."),
        ("AC-02", "US-02", "Given JWT hợp lệ và đủ tồn; When đặt hàng; Then API trả 201, snapshot giá, trừ tồn và tạo lịch sử kho trong cùng transaction."),
        ("AC-03", "US-02", "Given số lượng vượt tồn; When đặt hàng; Then trả 409, không sinh Order và tồn kho không thay đổi."),
        ("AC-04", "US-03", "Given Customer A; When truy cập orderId của Customer B; Then trả 403/404 và không lộ dữ liệu đơn."),
        ("AC-05", "US-04", "Given đơn Pending thuộc Customer; When hủy; Then trạng thái thành Cancelled và hoàn tồn nếu đã trừ theo chính sách."),
        ("AC-06", "US-05", "Given đơn Completed có sản phẩm; When gửi rating 1..5; Then tạo review; nếu chưa mua thì trả 403."),
        ("AC-07", "US-07", "Given Admin có inventory.import; When nhập số lượng dương; Then tồn tăng đúng và có transaction/audit log."),
        ("AC-08", "US-08", "Given đơn Confirmed; When chuyển Shipping bỏ qua Preparing; Then trả 409 và trạng thái không đổi."),
        ("AC-09", "US-09", "Given tài khoản bị khóa; When đăng nhập hoặc refresh; Then từ chối truy cập và ghi sự kiện bảo mật."),
        ("AC-10", "US-10", "Given khoảng ngày hợp lệ; When xem thống kê; Then tổng doanh thu chỉ tính đơn Completed và khớp chi tiết."),
    ]
    add_table(doc, ["ID", "US", "Given / When / Then"], criteria, [0.65, 0.6, 5.25], font_size=8.3)


def add_requirements(doc):
    add_heading(doc, "4. YÊU CẦU CHỨC NĂNG VÀ PHI CHỨC NĂNG", 1)
    add_heading(doc, "4.1. Functional Requirements (FR)", 2)
    frs = [
        ("FR-AUTH-01", "Cho phép đăng ký Customer bằng email duy nhất và mật khẩu đạt chính sách."),
        ("FR-AUTH-02", "Đăng nhập trả phiên xác thực; tài khoản khóa không được đăng nhập/refresh."),
        ("FR-AUTH-03", "Customer xem/cập nhật hồ sơ và đổi mật khẩu của chính mình."),
        ("FR-RBAC-01", "Backend ánh xạ User-Role-Permission và kiểm tra permission tại endpoint/service."),
        ("FR-RBAC-02", "Mặc định từ chối endpoint bảo vệ khi không có policy/permission phù hợp."),
        ("FR-CAT-01", "Guest xem danh mục và sản phẩm đang hoạt động."),
        ("FR-CAT-02", "Tìm, lọc, sắp xếp và phân trang sản phẩm."),
        ("FR-CAT-03", "Admin thêm/sửa/ẩn sản phẩm, ảnh và thông số kỹ thuật."),
        ("FR-CART-01", "Customer thêm/xóa/đổi số lượng và xem tổng tạm tính."),
        ("FR-ORDER-01", "Backend kiểm tra lại giá, trạng thái và tồn trước khi tạo đơn."),
        ("FR-ORDER-02", "OrderItem lưu tên/SKU/giá snapshot tại thời điểm đặt."),
        ("FR-ORDER-03", "Customer chỉ xem và hủy đơn thuộc chính mình khi trạng thái cho phép."),
        ("FR-ORDER-04", "Admin cập nhật trạng thái theo Pending→Confirmed→Preparing→Shipping→Completed hoặc Cancelled."),
        ("FR-INV-01", "Nhập/xuất kho tạo bản ghi lịch sử bất biến và cập nhật tồn trong transaction."),
        ("FR-INV-02", "Không cho phép tồn âm; cảnh báo khi tồn ≤ ngưỡng cấu hình."),
        ("FR-REV-01", "Chỉ Customer có OrderItem Completed được đánh giá 1-5 sao."),
        ("FR-REV-02", "Admin ẩn/hiện review và ghi lý do kiểm duyệt."),
        ("FR-REPORT-01", "Thống kê doanh thu, số đơn, top sản phẩm và tồn thấp theo ngày/tháng/năm."),
        ("FR-AUDIT-01", "Ghi audit cho thay đổi quyền, khóa tài khoản, kho, sản phẩm và trạng thái đơn."),
    ]
    add_table(doc, ["ID", "Yêu cầu"], frs, [1.2, 5.3], font_size=8.7)

    add_heading(doc, "4.2. Non-Functional Requirements (NFR)", 2)
    nfrs = [
        ("NFR-PERF-01", "Hiệu năng", "p95 GET catalog/chi tiết ≤ 2 giây với 10.000 sản phẩm, 50 người dùng đồng thời trên môi trường kiểm thử chuẩn."),
        ("NFR-PERF-02", "Hiệu năng", "p95 tạo đơn ≤ 3 giây; pageSize mặc định 20 và tối đa 100."),
        ("NFR-SEC-01", "Bảo mật", "100% endpoint quản trị yêu cầu JWT hợp lệ và permission; test 401/403/cross-user bắt buộc."),
        ("NFR-SEC-02", "Bảo mật", "Access token 15 phút; refresh token tối đa 7 ngày, xoay vòng và thu hồi khi đổi mật khẩu/khóa tài khoản."),
        ("NFR-DATA-01", "Toàn vẹn", "Tạo đơn và trừ kho là atomic; không tồn âm trong kiểm thử cạnh tranh."),
        ("NFR-REL-01", "Tin cậy", "Lỗi nghiệp vụ trả ProblemDetails có traceId; không trả stack trace ở production."),
        ("NFR-OPS-01", "Vận hành", "Có /health/live và /health/ready; structured log chứa timestamp, level, traceId và route."),
        ("NFR-MAINT-01", "Bảo trì", "Module nghiệp vụ tách rõ; service/domain quan trọng có unit test, mục tiêu coverage ≥ 70%."),
        ("NFR-UX-01", "Trải nghiệm", "Giao diện responsive từ 360px; trạng thái loading, empty, validation và lỗi API hiển thị rõ."),
        ("NFR-BACKUP-01", "Khôi phục", "Sao lưu PostgreSQL hằng ngày ở môi trường triển khai; diễn tập restore trước nghiệm thu."),
    ]
    add_table(doc, ["ID", "Nhóm", "Chỉ số kiểm chứng"], nfrs, [1.1, 1.0, 4.4], font_size=8.3)


def add_domain_data(doc, diagrams):
    add_heading(doc, "5. DOMAIN MODEL VÀ THIẾT KẾ DỮ LIỆU", 1)
    add_para(doc, "Mô hình dữ liệu tách định danh, phân quyền, catalog, giỏ hàng, đơn hàng, kho và đánh giá. Các snapshot trên đơn đảm bảo lịch sử không thay đổi khi tên hoặc giá sản phẩm được cập nhật sau này.")
    add_figure(doc, diagrams["erd"], "Hình 3. ERD rút gọn, cardinality và các khóa quan trọng", width=6.4)
    add_heading(doc, "5.1. Domain object/entity", 2)
    entities = [
        ("User, Role, Permission", "Danh tính, trạng thái, role và permission resource.action."),
        ("Category, Product", "Catalog, SKU, giá, tồn, trạng thái hiển thị và concurrency token."),
        ("Cart, CartItem", "Giỏ hiện hành; dữ liệu tạm tính, không phải nguồn giá cuối cùng."),
        ("Order, OrderItem", "Giao dịch mua và snapshot giá/tên/SKU/địa chỉ."),
        ("InventoryTransaction", "Sổ biến động IMPORT, EXPORT, ORDER, CANCEL_ADJUSTMENT; không sửa/xóa vật lý."),
        ("Payment", "COD hoặc SIMULATED; không lưu dữ liệu thẻ."),
        ("Review", "Rating, nội dung, trạng thái moderation và bằng chứng OrderItem."),
        ("AuditLog", "Ai, hành động, resource, resourceId, thời điểm, traceId và thay đổi quan trọng."),
    ]
    add_table(doc, ["Nhóm entity", "Trách nhiệm"], entities, [2.0, 4.5], font_size=9)

    add_heading(doc, "5.2. Constraint và index quan trọng", 2)
    constraints = [
        ("User", "UNIQUE(lower(email)); status ∈ ACTIVE, LOCKED", "idx_user_status"),
        ("Permission", "UNIQUE(code), mẫu resource.action", "uq_permission_code"),
        ("Product", "UNIQUE(sku); price ≥ 0; stock ≥ 0", "idx_product_category_active; GIN/trigram cho search khi cần"),
        ("CartItem", "UNIQUE(cart_id, product_id); quantity > 0", "idx_cartitem_cart"),
        ("Order", "status theo enum; total_amount ≥ 0", "idx_order_user_created; idx_order_status_created"),
        ("OrderItem", "quantity > 0; unit_price_snapshot ≥ 0", "idx_orderitem_order; idx_orderitem_product"),
        ("InventoryTransaction", "quantity_delta ≠ 0; lịch sử bất biến", "idx_inventory_product_created"),
        ("Review", "rating BETWEEN 1 AND 5; UNIQUE(user_id, order_item_id)", "idx_review_product_status"),
        ("RefreshToken", "UNIQUE(token_hash); expires_at > created_at", "idx_refresh_user_expires"),
        ("AuditLog", "append-only", "idx_audit_resource_created; idx_audit_actor_created"),
    ]
    add_table(doc, ["Bảng", "Constraint", "Index"], constraints, [1.2, 2.8, 2.5], font_size=8)


def add_system_design(doc, diagrams):
    add_heading(doc, "6. THIẾT KẾ HỆ THỐNG", 1)
    add_figure(doc, diagrams["container"], "Hình 4. Container và module của kiến trúc modular monolith")
    add_heading(doc, "6.1. Công nghệ mục tiêu", 2)
    add_table(
        doc,
        ["Lớp", "Công nghệ", "Trách nhiệm"],
        [
            ("Frontend", "React 19, Vite 8, Tailwind CSS 4, React Router 7, Axios", "Giao diện SPA, state UI, route, validation sơ bộ và gọi API."),
            ("Backend", "ASP.NET Core Web API .NET 10, EF Core", "Nghiệp vụ, validation quyết định, RBAC/JWT, transaction và API contract."),
            ("Database", "PostgreSQL", "Dữ liệu quan hệ, constraint, index, transaction và audit."),
            ("Tài liệu API", "OpenAPI/Swagger", "Mô tả endpoint, schema, status code và hỗ trợ kiểm thử."),
            ("UI prototype", "Stitch export chỉ làm tài liệu tham chiếu", "Không đưa demo_template vào production hoặc Git; triển khai lại thành component React."),
        ],
        [1.0, 2.2, 3.3],
        font_size=8.5,
    )
    doc.add_page_break()
    add_heading(doc, "6.2. Quy tắc phụ thuộc module", 2)
    for item in [
        "Controller chỉ nhận/trả DTO và gọi application service; không chứa truy vấn nghiệp vụ dài.",
        "Mỗi module sở hữu service, DTO, validation và repository/query liên quan; không truy cập vòng giữa module.",
        "Authorization kiểm tra tại endpoint và kiểm tra sở hữu resource tại service; frontend chỉ ẩn/hiện UI, không quyết định quyền.",
        "Transaction nằm ở use case cần atomicity, đặc biệt checkout, hủy đơn có hoàn tồn và nhập/xuất kho.",
        "Mọi lỗi chuẩn hóa thành application/problem+json với status, title, detail, instance và traceId.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6.3. Luồng chính", 2)
    add_figure(doc, diagrams["activity_order"], "Hình 5. Activity Diagram - đặt hàng và các nhánh lỗi", width=5.3)
    add_figure(doc, diagrams["sequence"], "Hình 6. Sequence Diagram - checkout với transaction giá/tồn kho")
    add_figure(doc, diagrams["activity_inventory"], "Hình 7. Activity Diagram - nhập kho và audit", width=5.3)


def add_api_contract(doc):
    add_heading(doc, "7. API CONTRACT", 1)
    add_para(doc, "Quy ước: base URL /api; JSON UTF-8; thời gian ISO-8601 UTC; tiền là decimal; phân trang trả items, page, pageSize, totalItems, totalPages. Lỗi dùng RFC 7807 ProblemDetails mở rộng traceId và errorCode.")
    api_rows = [
        ("POST", "/auth/register", "Public", "Tạo Customer", "201, 400, 409"),
        ("POST", "/auth/login", "Public", "Xác thực", "200, 400, 401, 423"),
        ("POST", "/auth/refresh", "Refresh cookie", "Xoay token", "200, 401"),
        ("POST", "/auth/logout", "Authenticated", "Thu hồi refresh token", "204, 401"),
        ("GET/PUT", "/me", "profile.self.*", "Xem/cập nhật hồ sơ", "200, 400, 401"),
        ("PUT", "/me/password", "Authenticated", "Đổi mật khẩu", "204, 400, 401"),
        ("GET", "/categories", "Public", "Danh mục hoạt động", "200"),
        ("GET", "/products", "Public", "Search/filter/sort/page", "200, 400"),
        ("GET", "/products/{id}", "Public", "Chi tiết sản phẩm", "200, 404"),
        ("POST/PUT", "/admin/products[/ {id}]", "product.create/update", "Tạo/sửa sản phẩm", "200/201, 400, 403, 409"),
        ("DELETE", "/admin/products/{id}", "product.delete", "Ẩn/xóa mềm", "204, 403, 404, 409"),
        ("GET/PUT", "/cart", "cart.read/update", "Lấy/cập nhật giỏ", "200, 400, 401, 409"),
        ("POST", "/orders", "order.create", "Checkout", "201, 400, 401, 403, 409"),
        ("GET", "/orders/me", "order.self.read", "Danh sách đơn của tôi", "200, 401"),
        ("GET", "/orders/me/{id}", "order.self.read + owner", "Chi tiết đơn", "200, 401, 403/404"),
        ("POST", "/orders/me/{id}/cancel", "order.self.cancel + owner", "Hủy đơn", "204, 401, 403, 409"),
        ("GET", "/admin/orders", "order.manage", "Danh sách đơn", "200, 403"),
        ("PATCH", "/admin/orders/{id}/status", "order.manage", "Chuyển trạng thái", "200, 403, 404, 409"),
        ("POST", "/admin/inventory/imports", "inventory.import", "Nhập kho", "201, 400, 403, 404"),
        ("POST", "/admin/inventory/exports", "inventory.export", "Xuất điều chỉnh", "201, 400, 403, 409"),
        ("GET", "/admin/inventory/transactions", "inventory.read", "Lịch sử kho", "200, 403"),
        ("POST", "/products/{id}/reviews", "review.create", "Đánh giá đã mua", "201, 400, 403, 409"),
        ("PATCH", "/admin/reviews/{id}/status", "review.moderate", "Kiểm duyệt", "200, 403, 404"),
        ("GET", "/admin/reports/overview", "report.read", "Dashboard theo khoảng ngày", "200, 400, 403"),
        ("PATCH", "/admin/users/{id}/status", "user.manage", "Khóa/mở tài khoản", "200, 403, 404, 409"),
    ]
    add_table(doc, ["Method", "Endpoint", "Quyền", "Mục đích", "Status"], api_rows, [0.65, 2.0, 1.35, 1.75, 0.75], font_size=7.1)
    add_heading(doc, "7.1. Ví dụ contract tạo đơn", 2)
    add_callout(doc, "Request", 'POST /api/orders { recipientName, phone, shippingAddress, paymentMethod, items:[{ productId, quantity }] }', fill=COLORS["lighter_blue"])
    add_callout(doc, "201 Created", '{ orderId, orderCode, status:"Pending", totalAmount, createdAt }', fill=COLORS["green_fill"], color=COLORS["green"])
    add_callout(doc, "409 Conflict", '{ type, title:"Business rule conflict", status:409, errorCode:"INSUFFICIENT_STOCK", detail, traceId }', fill=COLORS["red_fill"], color=COLORS["red"])


def add_state_rules(doc):
    add_heading(doc, "8. BUSINESS RULES VÀ STATE MACHINE", 1)
    rules = [
        ("BR-01", "Giá và tồn trên frontend chỉ mang tính hiển thị; backend luôn đọc lại khi checkout."),
        ("BR-02", "OrderItem lưu snapshot tên/SKU/giá; thay đổi Product không làm đổi đơn cũ."),
        ("BR-03", "Tồn kho không âm; cập nhật tồn và lịch sử kho trong cùng transaction."),
        ("BR-04", "Customer chỉ đọc/hủy đơn thuộc user_id của mình."),
        ("BR-05", "Chỉ đơn Pending được Customer hủy; Admin áp dụng chính sách hủy rõ ràng ở các trạng thái khác."),
        ("BR-06", "Review cần OrderItem thuộc đơn Completed và mỗi OrderItem chỉ có một review của người mua."),
        ("BR-07", "Xóa sản phẩm đã phát sinh giao dịch là xóa mềm/ẩn, không xóa vật lý."),
        ("BR-08", "Doanh thu chỉ tính đơn Completed; đơn Cancelled không tính."),
        ("BR-09", "Tài khoản LOCKED không thể login, refresh hoặc gọi endpoint được bảo vệ."),
        ("BR-10", "Mọi thao tác Admin nhạy cảm phải có audit log."),
    ]
    add_table(doc, ["ID", "Quy tắc"], rules, [0.75, 5.75], font_size=9)
    transitions = [
        ("Pending", "Confirmed", "Admin xác nhận"),
        ("Pending", "Cancelled", "Customer/Admin hủy"),
        ("Confirmed", "Preparing", "Admin bắt đầu chuẩn bị"),
        ("Confirmed", "Cancelled", "Admin hủy và hoàn tồn theo chính sách"),
        ("Preparing", "Shipping", "Bàn giao vận chuyển"),
        ("Shipping", "Completed", "Giao thành công"),
    ]
    add_heading(doc, "8.1. Chuyển trạng thái đơn hợp lệ", 2)
    add_table(doc, ["Từ", "Đến", "Điều kiện/hành động"], transitions, [1.1, 1.1, 4.3], font_size=9.2)


def add_adrs(doc):
    add_heading(doc, "9. TRADE-OFF VÀ ARCHITECTURE DECISION RECORDS", 1)
    adr_data = [
        ("ADR-001 - Modular monolith thay vì microservices",
         "Nhóm nhỏ, thời gian khoảng 6 tuần nhưng hệ thống có nhiều miền nghiệp vụ.",
         "A: một monolith không module; B: modular monolith; C: microservices.",
         "Chọn B. Một Web API triển khai đơn giản, transaction checkout dễ quản lý; module rõ để giảm coupling.",
         "Lợi ích: ít hạ tầng, debug và test end-to-end dễ. Chi phí: scale theo từng module còn hạn chế.",
         "Xem xét microservices khi có nhiều đội độc lập, >500 request/giây bền vững hoặc module báo cáo gây nghẽn riêng."),
        ("ADR-002 - PostgreSQL thay vì MySQL",
         "Cần transaction, constraint, index và truy vấn báo cáo; nhóm đã có kinh nghiệm cả PostgreSQL/MySQL.",
         "A: PostgreSQL; B: MySQL; C: SQLite.",
         "Chọn A vì constraint/transaction mạnh, kiểu dữ liệu và khả năng truy vấn báo cáo phù hợp; EF Core có provider ổn định.",
         "Chi phí: phải thống nhất provider/migration và không dùng cú pháp MySQL từ bài tuần 2.",
         "Xem xét lại nếu hạ tầng bắt buộc MySQL hoặc đội vận hành chỉ hỗ trợ MySQL."),
        ("ADR-003 - JWT ngắn hạn + refresh cookie HttpOnly",
         "SPA cần gọi API; giảm rủi ro token dài hạn bị lấy qua XSS.",
         "A: lưu JWT dài hạn ở localStorage; B: access JWT ngắn hạn trong memory + refresh HttpOnly; C: cookie session thuần.",
         "Chọn B để đáp ứng yêu cầu JWT của rubric và hạn chế JavaScript đọc refresh token.",
         "Chi phí: cần CORS credentials, CSRF protection cho refresh/logout và cơ chế rotation/revocation.",
         "Xem xét cookie session thuần nếu frontend/backend cùng site và không còn ràng buộc JWT."),
        ("ADR-004 - Transaction + row lock cho tồn kho",
         "Hai khách có thể đặt sản phẩm cuối cùng đồng thời.",
         "A: tin số tồn từ frontend; B: optimistic concurrency; C: transaction với khóa dòng khi checkout.",
         "Chọn C cho phiên bản đầu vì logic rõ và bảo vệ invariant stock ≥ 0; kết hợp timeout ngắn.",
         "Chi phí: contention khi sản phẩm rất nóng; transaction phải ngắn và index đúng.",
         "Xem xét reservation/queue khi tỷ lệ conflict hoặc thời gian chờ lock vượt ngưỡng vận hành."),
    ]
    for title, context, options, decision, consequences, trigger in adr_data:
        add_heading(doc, title, 2)
        add_para(doc, f"Context: {context}", bold_lead="Context:")
        add_para(doc, f"Options: {options}", bold_lead="Options:")
        add_para(doc, f"Decision & rationale: {decision}", bold_lead="Decision & rationale:")
        add_para(doc, f"Consequences: {consequences}", bold_lead="Consequences:")
        add_para(doc, f"Trigger xem xét lại: {trigger}", bold_lead="Trigger xem xét lại:")


def add_security(doc, diagrams):
    add_heading(doc, "10. THREAT MODEL, RBAC VÀ JWT", 1)
    add_figure(doc, diagrams["trust"], "Hình 8. Entry point, trust boundary và vùng dữ liệu", width=5.65)
    add_heading(doc, "10.1. Tài sản và mô hình đối thủ", 2)
    add_table(
        doc,
        ["Nhóm", "Nội dung"],
        [
            ("Asset", "Thông tin tài khoản/PII, password hash, token, giá/tồn, đơn hàng, quyền, báo cáo và audit log."),
            ("Entry point", "Form đăng nhập/đăng ký, search/query, checkout, upload ảnh, endpoint Admin, refresh cookie."),
            ("Attacker", "Guest ác ý, Customer cố truy cập chéo, tài khoản Admin bị chiếm, bot và người dùng nội bộ lạm quyền."),
            ("Trust boundary", "Browser↔Internet/API; API↔Database; API↔media storage; ranh giới Customer↔Admin."),
        ],
        [1.25, 5.25],
        font_size=9,
    )
    threats = [
        ("T-01", "Broken Object Level Authorization: đổi orderId để xem đơn người khác", "High", "Kiểm tra owner trong service; query theo orderId + currentUserId; test cross-user", "T-IDOR-01"),
        ("T-02", "Privilege escalation vào endpoint Admin", "High", "Permission policy, default deny, role/permission từ DB, audit", "T-RBAC-01/02"),
        ("T-03", "Brute force/credential stuffing", "High", "Rate limit, lockout, password policy, log sự kiện", "T-AUTH-03"),
        ("T-04", "Refresh/access token bị đánh cắp", "High", "Access 15 phút, refresh HttpOnly/Secure/SameSite, rotation, revoke", "T-AUTH-04"),
        ("T-05", "Overselling do race condition", "High", "Transaction + row lock, stock constraint và test đồng thời", "T-ORDER-03"),
        ("T-06", "Mass assignment sửa price/role/status", "High", "DTO allow-list; bỏ qua field server-owned; validation", "T-API-01"),
        ("T-07", "Stored/reflected XSS qua tên/review", "Medium", "React escaping, sanitize nội dung cho phép, CSP", "T-XSS-01"),
        ("T-08", "SQL injection", "Medium", "EF Core parameter hóa; không ghép raw SQL từ input", "T-SQL-01"),
        ("T-09", "CSRF tại refresh/logout dùng cookie", "Medium", "SameSite, anti-forgery/header, kiểm tra Origin", "T-CSRF-01"),
        ("T-10", "Upload ảnh độc hại/quá lớn", "Medium", "Allow-list MIME/đuôi, giới hạn kích thước, đổi tên, lưu ngoài web root", "T-UPLOAD-01"),
        ("T-11", "Lộ PII/token qua log", "High", "Redact header/body nhạy cảm; structured log allow-list", "T-LOG-01"),
        ("T-12", "Lạm dụng checkout/review", "Medium", "Rate limit, rule đã mua, idempotency key khi tạo đơn", "T-BIZ-01"),
    ]
    add_heading(doc, "10.2. Threat register", 2)
    add_table(doc, ["ID", "Threat", "Risk", "Mitigation", "Test"], threats, [0.48, 1.75, 0.52, 2.95, 0.8], font_size=6.9)

    add_heading(doc, "10.3. Ma trận Role-Permission", 2)
    permissions = [
        ("product.read / category.read / review.read", "✓", "✓", "✓"),
        ("auth.register / auth.login", "✓", "-", "-"),
        ("profile.self.* / cart.*", "-", "✓", "-"),
        ("order.create / order.self.read / order.self.cancel", "-", "✓", "-"),
        ("review.create", "-", "✓ có điều kiện đã mua", "-"),
        ("user.manage", "-", "-", "✓"),
        ("category.* / product.*", "-", "-", "✓"),
        ("inventory.read/import/export", "-", "-", "✓"),
        ("order.manage / review.moderate / report.read", "-", "-", "✓"),
    ]
    add_table(doc, ["Permission", "Guest", "Customer", "Admin"], permissions, [3.2, 0.85, 1.25, 1.2], font_size=8.2)
    add_callout(doc, "Nguyên tắc", "Backend default deny và least privilege. Ẩn nút trên frontend chỉ cải thiện trải nghiệm; mọi quyết định quyền và quyền sở hữu resource phải được kiểm tra lại trong API/service.")

    add_heading(doc, "10.4. JWT claims và vòng đời", 2)
    add_figure(doc, diagrams["jwt"], "Hình 9. JWT access token và refresh token flow")
    add_table(
        doc,
        ["Claim/thuộc tính", "Giá trị/quy tắc"],
        [
            ("sub", "user_id ổn định"),
            ("email", "Email chuẩn hóa; không dùng làm khóa sở hữu resource"),
            ("role", "Customer hoặc Admin; permission policy vẫn là kiểm tra chính"),
            ("permission", "Danh sách resource.action cần thiết hoặc nạp qua security stamp/cache ngắn"),
            ("jti", "Định danh access token phục vụ truy vết"),
            ("iss / aud", "Kiểm tra đúng issuer và audience"),
            ("iat / nbf / exp", "Access 15 phút; clock skew tối đa 1 phút"),
            ("Refresh token", "Random, lưu hash trong DB, HttpOnly Secure SameSite, 7 ngày, rotation và revoke"),
        ],
        [1.35, 5.15],
        font_size=8.8,
    )


def add_tests_traceability(doc):
    add_heading(doc, "11. KẾ HOẠCH KIỂM THỬ", 1)
    tests = [
        ("T-CAT-01", "Tìm/lọc/page catalog", "200; dữ liệu và total đúng; p95 đạt NFR"),
        ("T-AUTH-01", "Đăng nhập sai mật khẩu", "401, không cấp token"),
        ("T-AUTH-03", "Đăng nhập lặp vượt ngưỡng", "429/lockout và security log"),
        ("T-RBAC-01", "Customer gọi /admin/products", "403, không thay đổi dữ liệu"),
        ("T-RBAC-02", "Không JWT gọi endpoint bảo vệ", "401, không redirect HTML"),
        ("T-IDOR-01", "Customer A đọc đơn Customer B", "403/404, không lộ PII"),
        ("T-ORDER-01", "Đặt hàng đủ tồn", "201; Order/Items/InventoryTxn nhất quán"),
        ("T-ORDER-02", "Đặt vượt tồn", "409; rollback; stock không đổi"),
        ("T-ORDER-03", "Hai request tranh sản phẩm cuối", "Chỉ một request thành công; stock không âm"),
        ("T-ORDER-04", "Chuyển sai trạng thái", "409 và audit không ghi transition thành công"),
        ("T-REV-01", "Review khi chưa mua", "403"),
        ("T-REV-02", "Review sau đơn Completed", "201 và hiển thị sau moderation policy"),
        ("T-INV-01", "Nhập kho hợp lệ", "201; tồn tăng; có InventoryTxn/AuditLog"),
        ("T-API-01", "Body chứa price/role server-owned", "Bỏ qua hoặc 400; không mass assignment"),
        ("T-XSS-01", "Review chứa script", "Không thực thi trong UI; nội dung được xử lý"),
        ("T-PERF-01", "50 VU đọc catalog 10.000 SP", "p95 ≤ 2s, error rate <1%"),
        ("T-OPS-01", "DB không sẵn sàng", "/health/live vẫn live; /health/ready unhealthy"),
    ]
    add_table(doc, ["Test ID", "Tình huống", "Kết quả mong đợi"], tests, [0.9, 2.65, 2.95], font_size=8.1)

    add_heading(doc, "12. TRACEABILITY MATRIX", 1)
    rtm = [
        ("G-01", "UC-01", "US-01 / AC-01", "FR-CAT-01..03; NFR-PERF-01", "Category, Product", "GET /products", "Catalog", "T-CAT-01"),
        ("G-02", "UC-04", "US-02 / AC-02,03", "FR-ORDER-01,02; NFR-DATA-01", "Order, OrderItem, Product, InventoryTxn", "POST /orders", "Orders+Inventory", "T-ORDER-01..03"),
        ("G-03", "UC-05", "US-03,04 / AC-04,05", "FR-ORDER-03", "Order", "GET/POST /orders/me", "Orders", "T-IDOR-01"),
        ("G-03", "UC-06", "US-05 / AC-06", "FR-REV-01,02", "Review, OrderItem", "POST /products/{id}/reviews", "Reviews", "T-REV-01,02"),
        ("G-04", "UC-07", "US-06", "FR-CAT-03; FR-RBAC-01", "Product, Category", "/admin/products", "Catalog", "T-RBAC-01"),
        ("G-05", "UC-08", "US-07 / AC-07", "FR-INV-01,02; NFR-DATA-01", "Product, InventoryTxn, AuditLog", "/admin/inventory/*", "Inventory", "T-INV-01"),
        ("G-04", "UC-09", "US-08 / AC-08", "FR-ORDER-04", "Order, AuditLog", "PATCH /admin/orders/{id}/status", "Orders", "T-ORDER-04"),
        ("G-04", "UC-10", "US-09 / AC-09", "FR-AUTH-02; FR-RBAC-01,02", "User, Role, Permission", "PATCH /admin/users/{id}/status", "Identity", "T-AUTH-01,03"),
        ("G-06", "UC-11", "US-10 / AC-10", "FR-REPORT-01; NFR-PERF-01", "Order, OrderItem, Product", "GET /admin/reports/overview", "Reporting", "T-PERF-01"),
        ("G-04", "UC-02", "AC-04,09", "NFR-SEC-01,02", "RefreshToken, AuditLog", "/auth/*", "Identity/RBAC", "T-RBAC-02; T-AUTH-03"),
    ]
    add_table(doc, ["Goal", "UC", "US/AC", "FR/NFR", "Entity", "API", "Component", "Test"], rtm,
              [0.45, 0.45, 0.8, 1.25, 1.1, 1.15, 0.75, 0.55], font_size=6.3)


def add_operations(doc):
    add_heading(doc, "13. KHẢ NĂNG MỞ RỘNG VÀ VẬN HÀNH", 1)
    add_heading(doc, "13.1. Giả định tải và bottleneck", 2)
    add_table(
        doc,
        ["Giả định", "Thiết kế hiện tại", "Trigger nâng cấp"],
        [
            ("≤10.000 sản phẩm; 50 VU đọc", "Index category/status/price; pagination; projection DTO", "p95 catalog >2s trong 15 phút"),
            ("≤200 đơn/ngày trong đồ án", "Transaction ngắn + row lock", "Lock wait/conflict >2% checkout"),
            ("Dashboard theo ngày/tháng", "Query tổng hợp có index, cache in-memory 60s", "Báo cáo chiếm >30% DB CPU"),
            ("Ảnh sản phẩm vừa phải", "Lưu file ngoài DB, DB chỉ giữ URL/metadata", "Băng thông hoặc dung lượng vượt máy chủ"),
        ],
        [1.65, 2.75, 2.1],
        font_size=8.3,
    )
    add_heading(doc, "13.2. Error handling, health và log", 2)
    for item in [
        "Global exception handler trả ProblemDetails; lỗi validation 400, auth 401, forbidden 403, not found 404, conflict nghiệp vụ 409.",
        "Correlation/trace ID truyền từ frontend hoặc tạo ở API, xuất hiện trong response lỗi và structured log.",
        "Health checks tách live/ready; readiness kiểm tra PostgreSQL và dependency bắt buộc.",
        "Audit log append-only cho quyền, tài khoản, kho, sản phẩm và trạng thái đơn; không ghi password/token/PII không cần thiết.",
        "Backup hằng ngày; kiểm thử restore và ghi RPO/RTO thực tế trước khi triển khai thật.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "13.3. Technical debt và roadmap", 2)
    roadmap = [
        ("V1 - 6 tuần", "Modular monolith; COD/giả lập; một kho; search DB; UI responsive; đầy đủ RBAC/audit cốt lõi."),
        ("V1.1", "Email thông báo, tối ưu ảnh, cache catalog, cải thiện test tự động và dashboard."),
        ("V2", "Build PC/compatibility, voucher, cổng thanh toán và vận chuyển thật sau khi threat model/API được cập nhật."),
        ("Khi tải tăng", "Redis cache, background jobs, read model báo cáo hoặc tách service có số liệu chứng minh bottleneck."),
    ]
    add_table(doc, ["Mốc", "Nội dung"], roadmap, [1.0, 5.5], font_size=9)

    add_heading(doc, "13.4. Kế hoạch triển khai 6 tuần", 2)
    plan = [
        ("Tuần 1", "Chốt FR/NFR, ERD, API, ADR và RBAC; tạo migration nền."),
        ("Tuần 2", "Identity/RBAC/JWT; catalog và Admin sản phẩm/danh mục."),
        ("Tuần 3", "Frontend store: home, listing, detail, auth, cart."),
        ("Tuần 4", "Checkout, orders, state machine và inventory transaction."),
        ("Tuần 5", "Reviews, admin users/orders/inventory/reports; negative authorization tests."),
        ("Tuần 6", "Tích hợp, performance/security test, sửa lỗi, README, minh chứng và demo."),
    ]
    add_table(doc, ["Thời gian", "Kết quả cần đạt"], plan, [1.0, 5.5], font_size=9.2)


def add_ui_mapping(doc):
    add_heading(doc, "14. THIẾT KẾ GIAO DIỆN VÀ ÁNH XẠ CHỨC NĂNG", 1)
    add_para(doc, "Bộ giao diện Stitch là design reference. Khi triển khai, mỗi màn hình được dựng lại bằng React component, route và dữ liệu API; file HTML xuất từ Stitch không được dùng làm production source hoặc commit vào repository.")
    mapping = [
        ("Trang chủ", "/", "UC-01", "Header, CategoryNav, HeroBanner, ProductSection, Footer"),
        ("Danh sách sản phẩm", "/products", "UC-01", "FilterSidebar, SortBar, ProductGrid, Pagination"),
        ("Chi tiết sản phẩm", "/products/:id", "UC-01,03,06", "Gallery, PriceStock, Specs, Reviews, RelatedProducts"),
        ("Giỏ hàng", "/cart", "UC-03", "CartItem, QuantityControl, OrderSummary"),
        ("Thanh toán", "/checkout", "UC-04", "ShippingForm, PaymentSelector, CheckoutSummary"),
        ("Tài khoản/đơn", "/account/*", "UC-05,06", "ProfileForm, OrderList, OrderTimeline, ReviewForm"),
        ("Admin", "/admin/*", "UC-07..11", "AdminLayout, DataTable, StatusBadge, FormModal, Dashboard"),
    ]
    add_table(doc, ["Màn hình", "Route", "Use Case", "Component chính"], mapping, [1.2, 1.25, 0.9, 3.15], font_size=8.3)

    screenshot_root = Path(r"D:\HMNU\laptrinhwweb\website_Thuongmaidientu_linhkienmaytinh\demo_template\stitch_techhub_pc_e_commerce_interface")
    screenshots = [
        (screenshot_root / "techhub_pc_homepage_retail" / "screen.png", "Hình 10. Thiết kế tham chiếu trang chủ TechHub PC"),
        (screenshot_root / "techhub_pc_chi_ti_t_s_n_ph_m" / "screen.png", "Hình 11. Thiết kế tham chiếu trang chi tiết sản phẩm"),
        (screenshot_root / "techhub_pc_qu_n_tr_h_th_ng" / "screen.png", "Hình 12. Thiết kế tham chiếu dashboard quản trị"),
        (screenshot_root / "techhub_pc_qu_n_l_kho_h_ng_admin" / "screen.png", "Hình 13. Thiết kế tham chiếu quản lý kho"),
    ]
    for image_path, caption in screenshots:
        if image_path.exists():
            add_figure(doc, image_path, caption, width=6.25)


def add_conclusion_sources(doc):
    add_heading(doc, "15. KẾT LUẬN", 1)
    add_para(doc, "Thiết kế TechHub PC đã xác định rõ mục tiêu, tác nhân, luồng nghiệp vụ, dữ liệu, API, phân quyền, rủi ro và tiêu chí kiểm thử. Kiến trúc modular monolith cân bằng giữa khả năng hoàn thành trong phạm vi đồ án và khả năng mở rộng có điều kiện. Trọng tâm triển khai tiếp theo là hoàn thành frontend theo component/route, sau đó xây backend .NET 10 và kết nối PostgreSQL theo contract trong tài liệu.")
    add_callout(doc, "Điểm kiểm soát trước khi nộp", "Điền lớp, nhóm/sinh viên và giảng viên trên bìa; cập nhật mục lục trong Word; đổi tên file theo TenLop_TenNhom_TenDeTai; kiểm tra repository không chứa demo_template, .env, mật khẩu hoặc API key.", fill=COLORS["orange_fill"], color="7A3D00")

    add_heading(doc, "TÀI LIỆU THAM KHẢO VÀ CÔNG CỤ", 1)
    sources = [
        "Rubric_Phan_tich_Thiet_ke_HTTT.docx do giảng viên cung cấp.",
        "Website_Thuong_mai_Dien_tu_Linh_kien_May_tinh.docx - mô tả đề tài ban đầu.",
        "Microsoft Learn - ASP.NET Core 10 authentication và policy-based authorization: https://learn.microsoft.com/aspnet/core/security/authentication/ và https://learn.microsoft.com/aspnet/core/security/authorization/policies",
        "RFC 7519 - JSON Web Token (JWT): https://www.rfc-editor.org/rfc/rfc7519",
        "OWASP API Security Top 10 - 2023: https://owasp.org/API-Security/editions/2023/en/0x11-t10/",
        "PostgreSQL Documentation - Explicit Locking: https://www.postgresql.org/docs/current/explicit-locking.html",
        "Bộ thiết kế TechHub PC xuất từ Stitch, chỉ dùng làm reference cho UI.",
    ]
    for item in sources:
        add_bullet(doc, item)

    add_heading(doc, "Công cụ và công khai việc sử dụng AI", 2)
    add_para(doc, "Công cụ sử dụng gồm Microsoft Word/DOCX, React/Vite, ASP.NET Core, PostgreSQL, Swagger/Postman, Git/GitHub và Stitch. ChatGPT/Codex được dùng để hỗ trợ cấu trúc, rà soát tính nhất quán, sinh bản nháp tài liệu và kiểm tra trình bày. Người thực hiện chịu trách nhiệm đọc lại, hiệu chỉnh, triển khai, kiểm thử và giải thích toàn bộ quyết định trong buổi bảo vệ.")

    add_heading(doc, "PHỤ LỤC A - GLOSSARY", 1)
    glossary = [
        ("Authentication (AuthN)", "Xác định người dùng là ai."),
        ("Authorization (AuthZ)", "Xác định người dùng được phép làm gì trên resource cụ thể."),
        ("RBAC", "Phân quyền dựa trên role; trong thiết kế này role gom các permission resource.action."),
        ("JWT", "Token chứa claims được ký; không phải nơi lưu bí mật và phải kiểm tra issuer/audience/expiry."),
        ("Snapshot", "Bản sao tên/giá/địa chỉ lưu trên đơn để bảo toàn lịch sử."),
        ("Default deny", "Không có quyền rõ ràng thì từ chối truy cập."),
        ("Idempotency", "Gửi lại cùng yêu cầu không tạo hiệu ứng trùng lặp ngoài dự kiến."),
        ("ProblemDetails", "Định dạng lỗi HTTP chuẩn gồm type, title, status, detail, instance và trường mở rộng."),
    ]
    add_table(doc, ["Thuật ngữ", "Giải thích"], glossary, [1.6, 4.9], font_size=9)


def build():
    diagrams = generate_diagrams()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_document_control(doc)
    add_summary(doc)
    add_context_scope(doc, diagrams)
    add_actors_usecases(doc, diagrams)
    add_user_stories(doc)
    add_requirements(doc)
    add_domain_data(doc, diagrams)
    add_system_design(doc, diagrams)
    add_api_contract(doc)
    add_state_rules(doc)
    add_adrs(doc)
    add_security(doc, diagrams)
    add_tests_traceability(doc)
    add_operations(doc)
    add_ui_mapping(doc)
    add_conclusion_sources(doc)

    core = doc.core_properties
    core.title = "Phân tích và thiết kế hệ thống thông tin TechHub PC"
    core.subject = "Hồ sơ Use Case, Requirements, ERD, Architecture, RBAC/JWT, Threat Model và Traceability"
    core.author = "Nhóm thực hiện TechHub PC"
    core.keywords = "TechHub PC, ASP.NET Core, React, PostgreSQL, RBAC, JWT"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
